"""Формирование файла Excel по исходной форме (app/assets/template.xlsm)."""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Dict, List

from openpyxl import load_workbook

from .config import template_path
from .numberwords import amount_in_words, from_cents

SHEET_NAME = "Отчет к заполнению"

MONTHS_GENITIVE = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]

# Строки формы, соответствующие строкам в Excel
ADVANCE_ROWS = (26, 27, 28)
EXPENSE_ROWS = tuple(range(36, 46))


def _money(cents: int) -> float:
    return float(from_cents(cents))


def _clear_sheet_extras(workbook) -> None:
    """Оставляем в файле только бланк отчёта — без листов с примером и списками."""
    for name in ["Пример заполнения ", "Лист2", "Инструкция"]:
        if name in workbook.sheetnames:
            workbook.remove(workbook[name])
    try:
        names = list(workbook.defined_names.keys())
    except AttributeError:  # pragma: no cover - старые версии openpyxl
        return
    for name in names:
        if "Пример" in name:
            try:
                del workbook.defined_names[name]
            except (KeyError, TypeError):  # pragma: no cover
                pass


def build_report_xlsx(report: Dict[str, Any]) -> bytes:
    """Собирает заполненный бланк отчёта в формате .xlsx."""
    workbook = load_workbook(template_path(), keep_vba=False, data_only=False)
    if SHEET_NAME not in workbook.sheetnames:
        raise RuntimeError("В шаблоне отсутствует лист «Отчет к заполнению»")
    sheet = workbook[SHEET_NAME]

    report_date = date.fromisoformat(str(report["report_date"])[:10])
    advances: List[dict] = list(report.get("advances") or [])
    expenses: List[dict] = list(report.get("expenses") or [])
    totals = report.get("totals") or {}
    approved_cents = int(totals.get("approved_cents") or 0)

    # --- Шапка отчёта -----------------------------------------------------
    sheet["F16"] = report_date.day
    sheet["H16"] = MONTHS_GENITIVE[report_date.month - 1]
    sheet["L16"] = report_date.year

    sheet["E18"] = report.get("person_name", "")
    sheet["E20"] = report.get("person_position", "")
    sheet["E22"] = report.get("purpose", "")

    # --- Полученные авансы (до трёх строк) --------------------------------
    for index, row in enumerate(ADVANCE_ROWS):
        if index < len(advances):
            advance = advances[index]
            raw_date = advance.get("date")
            if raw_date:
                try:
                    sheet[f"A{row}"] = datetime.fromisoformat(str(raw_date)[:10])
                except ValueError:
                    sheet[f"A{row}"] = str(raw_date)
            sheet[f"M{row}"] = _money(int(advance.get("amount_cents") or 0))
        else:
            sheet[f"A{row}"] = None
            sheet[f"M{row}"] = None

    # --- Расходы (десять строк) ------------------------------------------
    for index, row in enumerate(EXPENSE_ROWS):
        if index < len(expenses):
            expense = expenses[index]
            sheet[f"A{row}"] = expense.get("date", "")
            sheet[f"C{row}"] = expense.get("document", "")
            sheet[f"E{row}"] = expense.get("document_number", "")
            sheet[f"H{row}"] = expense.get("purpose", "")
            sheet[f"P{row}"] = _money(int(expense.get("amount_cents") or 0))
        else:
            for column in ("A", "C", "E", "H", "P"):
                sheet[f"{column}{row}"] = None

    # --- Итоговые значения: пишем числа, а не формулы ---------------------
    # (чтобы файл корректно открывался в любом просмотрщике без пересчёта)
    sheet["M29"] = _money(int(totals.get("received_cents") or 0))
    sheet["M30"] = _money(int(totals.get("spent_cents") or 0))
    sheet["M31"] = _money(int(totals.get("balance_cents") or 0))
    sheet["P46"] = _money(int(totals.get("spent_cents") or 0))
    sheet["F52"] = _money(approved_cents)

    words = amount_in_words(approved_cents)
    sheet["S4"] = words
    sheet["S5"] = words
    sheet["I6"] = words
    sheet["K49"] = report.get("person_name", "")
    sheet["N10"] = report.get("director_name", "") or sheet["N10"].value

    _clear_sheet_extras(workbook)
    try:
        workbook.calculation.fullCalcOnLoad = True
    except AttributeError:  # pragma: no cover
        pass

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def report_filename(report: Dict[str, Any]) -> str:
    number = str(report.get("number") or "report")
    person = str(report.get("person_name") or "").strip().replace(" ", "_")
    filename = f"Авансовый_отчет_{number}"
    if person:
        filename += f"_{person}"
    return filename + ".xlsx"
